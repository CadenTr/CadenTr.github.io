"""Build the Selected Work grid in the portfolio resume."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REPO = Path(__file__).resolve().parents[1]
SOURCE = REPO / "assets" / "documents" / "caden-trahan-resume.docx"


CONTENT = (
    (
        (
            "Sigsbee Utility Corridor Exhibits",
            "Civil 3D · Bluebeam · GIS",
            "Drafted ~45 exhibits covering 21 parcels",
            "Set up DWGs, title blocks, points, and parcel linework",
            "Applied deed research and PLS review comments",
        ),
        (
            "Blankenship Phase 2B Plat",
            "Civil 3D · Bluebeam",
            "Drafted a four-sheet, 21-lot subdivision plat",
            "Mapped ROW, open space, BMP, and easement areas",
            "Checked 4,431.23 ft boundary with 0.03 ft closure",
        ),
    ),
    (
        (
            "BMW ALVA Pond As-Built",
            "Civil 3D · Excel",
            "Calculated stage storage from as-built pond data",
            "Documented 14.173 acre-ft of total storage",
            "Checked embankment and littoral shelf information",
        ),
        (
            "Field-to-Finish Physical Survey",
            "Leica Total Station · Civil 3D · Bluebeam · Polaris",
            "Operated a Leica total station in the field",
            "Located control, boundary, utilities, and site features",
            "Processed observations and drafted the final survey",
        ),
    ),
    (
        (
            "Utility As-Builts",
            "Civil 3D · Bluebeam",
            "Drafted storm, sewer, and water as-builts",
            "Checked rims, inverts, slopes, and linework",
            "Added structure callouts and sheet revisions",
        ),
        (
            "As-Built Production and Coordination",
            "Civil 3D · Bluebeam",
            "Managed scheduling, communication, and drafting",
            "Coordinated reshoots for missing or conflicting data",
            "Prepared final submittals for county recordation",
        ),
    ),
)


def clear_paragraphs(cell) -> None:
    tc = cell._tc
    for paragraph in list(cell.paragraphs):
        tc.remove(paragraph._p)


def set_cell_margins(cell, top=38, start=70, bottom=38, end=70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_line(cell, text: str, index: int) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1.0 if index < 2 else 0.4)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = index < 4

    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(9.15 if index < 2 else 8.75)
    run.bold = index == 0
    run.italic = index == 1


def main() -> None:
    document = Document(SOURCE)
    table = document.tables[0]

    for row_index, row in enumerate(table.rows):
        row.height = Inches(0.88)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            clear_paragraphs(cell)
            for line_index, line in enumerate(CONTENT[row_index][cell_index]):
                add_line(cell, line, line_index)

    document.save(SOURCE)
    print(SOURCE)


if __name__ == "__main__":
    main()
